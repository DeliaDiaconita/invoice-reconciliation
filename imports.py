import csv
from openpyxl import load_workbook
import io


# CITIRE PLATI DIN CSV
def load_payments(file):
    payments = []

    text_file = io.StringIO(file.getvalue().decode("utf-8-sig"))
    reader = csv.DictReader(text_file, delimiter=";")

    for row in reader:
        amount = float(row["suma"])
        if amount > 0:
            payment = {
                "payment_id": len(payments) + 1,
                "name": row["nume beneficiar/ordonator"],
                "description": row["detalii tranzactie"],
                "amount": amount,
                "payment_date": row["data procesarii"],
            }

            payments.append(payment)
    return payments


# CITIRE FACTURI DIN EXCEL
def load_invoices(file):
    workbook = load_workbook(file)
    sheet = workbook.active

    invoices = []

    for row in range(5, sheet.max_row + 1):
        client = sheet[f"B{row}"].value
        invoice_number = sheet[f"E{row}"].value
        issue_date = sheet[f"F{row}"].value
        amount = sheet[f"J{row}"].value

        if invoice_number is None:
            continue

        invoice = {
            "client": client,
            "invoice_number": invoice_number,
            "issue_date": issue_date.replace("/", "."),
            "amount": amount,
        }
        invoices.append(invoice)
    return invoices
